"""DOCX fixtures are generated here so source tests remain reviewable text."""

from __future__ import annotations

import zipfile

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from bn_en_translate.book.formats.docx import (
    DocxExportError,
    DocxExportWarning,
    DocxImportError,
    DocxImportWarning,
    DocxReader,
    DocxWriter,
    validate_docx_round_trip,
)
from bn_en_translate.book.schema import BlockKind, BookBlock, BookDocument


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _write_fixture(path) -> None:
    document = Document()
    document.core_properties.title = "বাংলা উপন্যাস"
    document.core_properties.author = "লেখক"
    document.core_properties.subject = "fixture"
    document.add_heading("প্রথম অধ্যায়", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("গাঢ় ").bold = True
    paragraph.add_run("তির্যক ").italic = True
    paragraph.add_run("দাগ").underline = True
    _add_hyperlink(paragraph, " লিংক", "https://example.test/chapter")
    document.add_paragraph("তালিকা", style="List Bullet")
    document.add_paragraph("উদ্ধৃতি", style="Quote")
    document.add_paragraph("")
    document.add_paragraph("***")
    document.add_heading("দ্বিতীয় অধ্যায়", level=1)
    document.add_heading("উপশিরোনাম", level=2)
    document.add_paragraph("শেষ।")
    document.save(path)


def test_docx_import_preserves_structure_runs_and_metadata(tmp_path) -> None:
    source = tmp_path / "book.bn.docx"
    _write_fixture(source)

    document = DocxReader().read(source)

    assert len(document.chapters) == 2
    assert [block.kind for block in document.blocks] == [
        BlockKind.CHAPTER_HEADING,
        BlockKind.PARAGRAPH,
        BlockKind.LIST_ITEM,
        BlockKind.BLOCK_QUOTE,
        BlockKind.BLANK,
        BlockKind.SCENE_BREAK,
        BlockKind.CHAPTER_HEADING,
        BlockKind.HEADING,
        BlockKind.PARAGRAPH,
    ]
    assert [block.block_id for block in document.blocks] == [
        "c0001-b000001",
        "c0001-b000002",
        "c0001-b000003",
        "c0001-b000004",
        "c0001-b000005",
        "c0001-b000006",
        "c0002-b000007",
        "c0002-b000008",
        "c0002-b000009",
    ]
    assert document.chapters[0].title == "প্রথম অধ্যায়"
    assert document.chapters[1].title == "দ্বিতীয় অধ্যায়"
    assert document.metadata.source_format == "docx"
    assert document.metadata.title == "বাংলা উপন্যাস"
    assert document.metadata.core_properties["author"] == "লেখক"
    assert document.metadata.core_properties["subject"] == "fixture"

    mixed_runs = document.blocks[1].runs
    assert [(run.text, run.bold, run.italic, run.underline, run.href) for run in mixed_runs] == [
        ("গাঢ় ", True, False, False, None),
        ("তির্যক ", False, True, False, None),
        ("দাগ", False, False, True, None),
        (" লিংক", False, False, False, "https://example.test/chapter"),
    ]
    assert document.blocks[2].attrs["paragraph_style"] == "List Bullet"
    assert document.blocks[2].attrs["list"]["level"] == 0
    assert document.blocks[0].attrs["source_locator"] == "word/document.xml:p[1]"


def test_docx_import_is_deterministic_for_same_source(tmp_path) -> None:
    source = tmp_path / "book.bn.docx"
    _write_fixture(source)

    first = DocxReader().read(source)
    second = DocxReader().read(source)

    assert first.document_id == second.document_id
    assert [(block.block_id, block.source_hash) for block in first.blocks] == [
        (block.block_id, block.source_hash) for block in second.blocks
    ]


def test_docx_import_warns_for_unsupported_tables(tmp_path) -> None:
    source = tmp_path / "table.docx"
    document = Document()
    document.add_paragraph("পাঠ্য")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "সারণি"
    document.save(source)

    with pytest.warns(DocxImportWarning, match="tables"):
        imported = DocxReader().read(source)
    assert [block.source_text for block in imported.blocks] == ["পাঠ্য"]


def test_docx_import_retains_note_reference_and_warns(tmp_path) -> None:
    source = tmp_path / "notes.docx"
    document = Document()
    paragraph = document.add_paragraph("পাঠ্য")
    reference = OxmlElement("w:footnoteReference")
    reference.set(qn("w:id"), "7")
    paragraph.add_run()._r.append(reference)
    document.save(source)

    with pytest.warns(DocxImportWarning, match="footnote/endnote"):
        imported = DocxReader().read(source)
    assert imported.blocks[0].attrs["note_references"] == ({"kind": "footnote", "id": "7"},)


def test_docx_import_rejects_non_package_input(tmp_path) -> None:
    source = tmp_path / "locked.docx"
    source.write_bytes(b"not a DOCX")

    with pytest.raises(DocxImportError, match="password-protected, corrupt"):
        DocxReader().read(source)
    assert not zipfile.is_zipfile(source)


def test_docx_export_preserves_semantic_coverage_and_private_ids(tmp_path) -> None:
    source = tmp_path / "book.bn.docx"
    destination = tmp_path / "book.en.docx"
    _write_fixture(source)
    imported = DocxReader().read(source)
    translations = {
        block.block_id: (
            block.source_text if block.ordinal == 2 else f"Target {block.ordinal}"
        )
        for block in imported.blocks
    }

    with pytest.warns(DocxExportWarning, match="inline styles"):
        DocxWriter().write(imported, translations, destination)

    report = validate_docx_round_trip(imported, destination)
    assert report.ok
    output = Document(destination)
    assert all(
        block.block_id not in paragraph.text
        for paragraph in output.paragraphs
        for block in imported.blocks
    )
    assert output.paragraphs[0].style.name.startswith("Heading 1")
    with zipfile.ZipFile(destination) as package:
        metadata = package.read("customXml/bn_book_blocks.json").decode("utf-8")
    assert imported.blocks[0].block_id in metadata


def test_docx_export_does_not_guess_changed_inline_run_alignment(tmp_path) -> None:
    source = tmp_path / "book.bn.docx"
    destination = tmp_path / "book.en.docx"
    _write_fixture(source)
    imported = DocxReader().read(source)
    body = imported.blocks[1]
    # Equal character counts still do not make Bengali/English run offsets
    # reliable, so changed text must remain one unstyled target run.
    target = "A" * len(body.source_text)
    writer = DocxWriter()
    with pytest.warns(DocxExportWarning, match="inline styles"):
        writer.write(
            imported,
            {
                block.block_id: target if block.block_id == body.block_id else block.source_text
                for block in imported.blocks
            },
            destination,
        )
    assert writer.findings[0]["rule"] == "inline_style_projection"
    exported = DocxReader().read(destination)
    runs = exported.blocks[1].runs
    assert len(runs) == 1
    assert not runs[0].bold and not runs[0].italic and not runs[0].underline


def test_docx_export_reconstructs_direct_numbering_metadata(tmp_path) -> None:
    source = tmp_path / "book.bn.docx"
    destination = tmp_path / "book.en.docx"
    _write_fixture(source)
    imported = DocxReader().read(source)
    list_block = imported.blocks[2]
    attrs = dict(list_block.attrs)
    attrs["paragraph_style"] = "Normal"
    attrs["list"] = {"level": 1, "num_id": "1"}
    replacement = BookBlock.create(
        block_id=list_block.block_id,
        chapter_id=list_block.chapter_id,
        ordinal=list_block.ordinal,
        kind=BlockKind.LIST_ITEM,
        source_text=list_block.source_text,
        attrs=attrs,
    )
    blocks = (*imported.blocks[:2], replacement, *imported.blocks[3:])
    modified = BookDocument(
        document_id=imported.document_id,
        metadata=imported.metadata,
        chapters=imported.chapters,
        blocks=blocks,
    )
    modified.validate()
    DocxWriter().write(
        modified,
        {block.block_id: block.source_text for block in modified.blocks},
        destination,
    )
    exported = DocxReader().read(destination)
    assert exported.blocks[2].kind is BlockKind.LIST_ITEM
    assert exported.blocks[2].attrs["list"]["level"] == 1


def test_docx_export_is_atomic_when_translation_is_missing(tmp_path) -> None:
    source = tmp_path / "book.bn.docx"
    destination = tmp_path / "book.en.docx"
    _write_fixture(source)
    destination.write_bytes(b"existing output")
    imported = DocxReader().read(source)
    with pytest.raises(DocxExportError, match="missing translation"):
        DocxWriter().write(imported, {}, destination)
    assert destination.read_bytes() == b"existing output"
